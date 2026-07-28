"""Build deterministic, synthetic evaluation documents.

The generated material is original project content and safe to redistribute.
Run this script only when the committed fixtures need to be regenerated.
"""

from __future__ import annotations

from pathlib import Path

import docx
from pptx import Presentation
from pptx.util import Inches


FIXTURE_DIR = Path(__file__).parent / "fixtures"


def _pdf_escape(value: str) -> str:
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _write_simple_pdf(path: Path, pages: list[list[str]]) -> None:
    """Write a small text-only PDF using the standard Helvetica font."""

    font_id = 3 + len(pages) * 2
    objects: list[bytes] = []
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{3 + index * 2} 0 R" for index in range(len(pages)))
    objects.append(
        f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode()
    )
    for index, lines in enumerate(pages):
        page_id = 3 + index * 2
        content_id = page_id + 1
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 {font_id} 0 R >> >> "
                f"/Contents {content_id} 0 R >>"
            ).encode()
        )
        commands = ["BT", "/F1 12 Tf", "72 730 Td", "15 TL"]
        for line_index, line in enumerate(lines):
            if line_index:
                commands.append("T*")
            commands.append(f"({_pdf_escape(line)}) Tj")
        commands.append("ET")
        stream = "\n".join(commands).encode("ascii")
        objects.append(
            f"<< /Length {len(stream)} >>\nstream\n".encode()
            + stream
            + b"\nendstream"
        )
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for object_id, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{object_id} 0 obj\n".encode())
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode()
    )
    path.write_bytes(output)


def build_pdf() -> None:
    pages = [
        [
            "Algorithms Notes - Search",
            "Linear search examines items sequentially and takes O(n) time in the worst case.",
            "Binary search requires sorted input and halves the remaining interval each step.",
        ],
        [
            "Sorting",
            "Merge sort divides the input, sorts each half, and merges the results.",
            "Its worst-case running time is O(n log n) and it needs O(n) auxiliary space.",
        ],
        [
            "Graphs",
            "Breadth-first search uses a queue and finds shortest paths in unweighted graphs.",
            "Depth-first search uses a stack or recursion and is useful for cycle detection.",
        ],
        [
            "Dynamic Programming",
            "Dynamic programming stores solutions to overlapping subproblems.",
            "Memoization is top-down while tabulation is bottom-up.",
        ],
        [
            "Hashing",
            "A hash table provides expected O(1) lookup with a suitable hash function.",
            "Separate chaining stores colliding keys in per-bucket collections.",
        ],
    ]
    _write_simple_pdf(FIXTURE_DIR / "algorithms_notes.pdf", pages)


def build_pptx() -> None:
    slides = [
        (
            "Relational Foundations",
            "A primary key uniquely identifies each row. A foreign key references a key in another table.",
        ),
        (
            "Normalization",
            "Third normal form removes transitive dependencies on non-key attributes.",
        ),
        (
            "Transactions",
            "ACID means atomicity, consistency, isolation, and durability.",
        ),
        (
            "Indexes",
            "A B-tree index supports logarithmic search and efficient range queries.",
        ),
        (
            "Joins",
            "An inner join returns matching rows. A left join also retains unmatched rows from the left table.",
        ),
    ]
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    for title, body in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    presentation.save(FIXTURE_DIR / "database_slides.pptx")


def build_docx() -> None:
    paragraphs = [
        "The application layer provides services such as HTTP and DNS to user programs.",
        "HTTP commonly uses a request-response interaction between a client and server.",
        "The transport layer provides process-to-process delivery.",
        "TCP is connection-oriented and offers reliable ordered byte delivery.",
        "The network layer routes packets between hosts across interconnected networks.",
        "IP provides best-effort datagram delivery without a reliability guarantee.",
        "The link layer transfers frames across one local link.",
        "Ethernet switches forward frames using learned MAC address tables.",
        "DNS translates domain names into IP addresses through a distributed hierarchy.",
        "A recursive resolver contacts other DNS servers on behalf of a client.",
    ]
    document = docx.Document()
    document.add_heading("Networking Handout", level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(FIXTURE_DIR / "networking_handout.docx")


def build_txt() -> None:
    lines = [
        "Study Design and Statistics",
        "A population is the complete group a study aims to understand.",
        "A sample is the subset of the population that is actually observed.",
        "Random sampling reduces selection bias in estimating population properties.",
        "A parameter describes a population.",
        "A statistic describes a sample.",
        "The mean is the arithmetic average.",
        "The median is the middle ordered value.",
        "The standard deviation measures spread around the mean.",
        "An outlier can influence the mean more strongly than the median.",
        "Correlation measures association but does not establish causation.",
        "A confounder is related to both an exposure and an outcome.",
        "Random assignment helps balance confounders between treatment groups.",
        "Blinding reduces behavior and measurement changes caused by knowing assignments.",
        "A control group provides a comparison for the treatment group.",
        "A null hypothesis represents a baseline claim.",
        "A p-value is computed assuming the null hypothesis is true.",
        "Statistical significance does not necessarily imply practical importance.",
        "Confidence intervals communicate estimate uncertainty.",
        "A 95 percent confidence procedure captures the true parameter in 95 percent of repeated samples.",
        "Machine Learning Study Guide",
        "Supervised learning uses labeled examples.",
        "Unsupervised learning seeks structure in unlabeled data.",
        "Classification predicts categories while regression predicts numeric values.",
        "Training data is used to fit model parameters.",
        "Validation data supports model and hyperparameter selection.",
        "Test data estimates performance after model choices are fixed.",
        "Overfitting occurs when a model learns noise and fails to generalize.",
        "Regularization discourages overly complex fitted models.",
        "Cross-validation rotates held-out folds to estimate generalization.",
        "Precision is the fraction of predicted positives that are correct.",
        "Recall is the fraction of actual positives that are found.",
        "The F1 score is the harmonic mean of precision and recall.",
        "Accuracy can mislead on severely imbalanced classes.",
        "A confusion matrix counts predicted and actual class combinations.",
        "Feature scaling can matter for distance-based algorithms.",
        "Gradient descent updates parameters opposite the loss gradient.",
        "A learning rate controls the size of each gradient update.",
        "Data leakage exposes information unavailable at real prediction time.",
        "A reproducible experiment records data, code, parameters, and random seeds.",
    ]
    (FIXTURE_DIR / "study_guide.txt").write_text(
        "\n".join(lines) + "\n",
        encoding="utf-8",
    )


def main() -> None:
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    build_pdf()
    build_pptx()
    build_docx()
    build_txt()
    print(f"Built evaluation fixtures in {FIXTURE_DIR}")


if __name__ == "__main__":
    main()

