from __future__ import annotations

import io

import docx
import pytest
from pptx import Presentation
from pptx.util import Inches

from ingest import (
    DocumentExtractionError,
    EmptyDocumentError,
    UnsupportedFileTypeError,
    compute_file_hash,
    extract_document,
)


def test_file_hash_is_stable_and_content_sensitive():
    assert compute_file_hash(b"abc") == compute_file_hash(b"abc")
    assert compute_file_hash(b"abc") != compute_file_hash(b"abcd")


def test_txt_encoding_and_line_locations():
    text = "caf\xe9\nsecond line\nthird line".encode("windows-1252")
    units = extract_document(text, "notes.txt", "Reading Materials")
    assert units[0].source == "notes.txt"
    assert units[0].location == "Lines 1-3"
    assert "café" in units[0].text


def test_docx_preserves_paragraph_positions():
    document = docx.Document()
    document.add_paragraph("First paragraph")
    document.add_paragraph("")
    document.add_paragraph("Third paragraph")
    buffer = io.BytesIO()
    document.save(buffer)
    units = extract_document(buffer.getvalue(), "notes.docx", "Reading Materials")
    assert [unit.location for unit in units] == ["Paragraph 1", "Paragraph 3"]


def test_pptx_preserves_slide_positions():
    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    slide1 = presentation.slides.add_slide(blank)
    box = slide1.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text = "First slide evidence"
    presentation.slides.add_slide(blank)
    slide3 = presentation.slides.add_slide(blank)
    box = slide3.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text = "Third slide evidence"
    buffer = io.BytesIO()
    presentation.save(buffer)
    units = extract_document(buffer.getvalue(), "slides.pptx", "Reading Materials")
    assert [unit.location for unit in units] == ["Slide 1", "Slide 3"]


def test_empty_file_is_rejected():
    with pytest.raises(EmptyDocumentError):
        extract_document(b"", "empty.txt", "Reading Materials")


def test_unsupported_extension_is_rejected():
    with pytest.raises(UnsupportedFileTypeError):
        extract_document(b"content", "notes.csv", "Reading Materials")


def test_textless_docx_is_rejected():
    document = docx.Document()
    buffer = io.BytesIO()
    document.save(buffer)
    with pytest.raises(DocumentExtractionError):
        extract_document(buffer.getvalue(), "empty.docx", "Reading Materials")

