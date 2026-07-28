"""Source-aware extraction and deterministic token-based chunking."""

from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass
from pathlib import Path

import chardet
import docx
import tiktoken
from pptx import Presentation
from PyPDF2 import PdfReader


class IngestionError(Exception):
    """Base exception for expected ingestion failures."""


class UnsupportedFileTypeError(IngestionError):
    """Raised when an uploaded extension is not supported."""


class EmptyDocumentError(IngestionError):
    """Raised when an uploaded file has no bytes or usable content."""


class DocumentExtractionError(IngestionError):
    """Raised when a supported file cannot yield usable text."""


@dataclass(frozen=True)
class SourceUnit:
    text: str
    source: str
    location: str
    category: str
    unit_index: int


@dataclass
class DocumentChunk:
    chunk_id: str
    text: str
    source: str
    location: str
    category: str
    unit_index: int
    chunk_index: int
    token_count: int
    embedding: list[float] | None = None


_ENCODING = tiktoken.get_encoding("cl100k_base")
_SUPPORTED_EXTENSIONS = {".pdf", ".pptx", ".docx", ".txt"}


def encode_text(text: str) -> list[int]:
    """Encode text with a deterministic general-purpose OpenAI tokenizer."""

    return _ENCODING.encode(text)


def decode_tokens(tokens: list[int]) -> str:
    """Decode tokens emitted by :func:`encode_text`."""

    return _ENCODING.decode(tokens)


def count_tokens(text: str) -> int:
    """Return an exact tokenizer count for the configured encoding."""

    return len(encode_text(text))


def compute_file_hash(file_bytes: bytes) -> str:
    """Return a SHA-256 hash for file identity and duplicate detection."""

    return hashlib.sha256(file_bytes).hexdigest()


def _clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_pdf(file_bytes: bytes, filename: str, category: str) -> list[SourceUnit]:
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        units = []
        for index, page in enumerate(reader.pages):
            text = _clean_text(page.extract_text() or "")
            if text:
                units.append(
                    SourceUnit(text, filename, f"Page {index + 1}", category, index)
                )
    except Exception as exc:
        raise DocumentExtractionError(f"Could not read PDF '{filename}'.") from exc
    if not units:
        raise DocumentExtractionError(
            "The PDF contains no extractable text. Image-only PDFs require OCR, "
            "which this application intentionally does not provide."
        )
    return units


def _extract_pptx(
    file_bytes: bytes, filename: str, category: str
) -> list[SourceUnit]:
    try:
        presentation = Presentation(io.BytesIO(file_bytes))
        units = []
        for index, slide in enumerate(presentation.slides):
            parts = []
            for shape in slide.shapes:
                text = _clean_text(getattr(shape, "text", "") or "")
                if text:
                    parts.append(text)
            if parts:
                units.append(
                    SourceUnit(
                        "\n".join(parts),
                        filename,
                        f"Slide {index + 1}",
                        category,
                        index,
                    )
                )
    except Exception as exc:
        raise DocumentExtractionError(f"Could not read PPTX '{filename}'.") from exc
    if not units:
        raise DocumentExtractionError(
            "The presentation contains no extractable text. Text inside images is "
            "not supported."
        )
    return units


def _extract_docx(
    file_bytes: bytes, filename: str, category: str
) -> list[SourceUnit]:
    try:
        document = docx.Document(io.BytesIO(file_bytes))
        units = []
        for index, paragraph in enumerate(document.paragraphs):
            text = _clean_text(paragraph.text)
            if text:
                units.append(
                    SourceUnit(
                        text,
                        filename,
                        f"Paragraph {index + 1}",
                        category,
                        index,
                    )
                )
    except Exception as exc:
        raise DocumentExtractionError(f"Could not read DOCX '{filename}'.") from exc
    if not units:
        raise DocumentExtractionError("The DOCX contains no usable paragraph text.")
    return units


def _decode_txt(file_bytes: bytes) -> str:
    detected = chardet.detect(file_bytes)
    candidates = [detected.get("encoding"), "utf-8-sig", "utf-8", "windows-1252"]
    for encoding in dict.fromkeys(item for item in candidates if item):
        try:
            return file_bytes.decode(encoding)
        except (LookupError, UnicodeDecodeError):
            continue
    return file_bytes.decode("utf-8", errors="replace")


def _extract_txt(file_bytes: bytes, filename: str, category: str) -> list[SourceUnit]:
    text = _decode_txt(file_bytes)
    lines = text.splitlines()
    units: list[SourceUnit] = []
    start: int | None = None
    buffered: list[str] = []

    def emit(end_line: int) -> None:
        nonlocal start, buffered
        cleaned = _clean_text("\n".join(buffered))
        if cleaned and start is not None:
            location = (
                f"Line {start}" if start == end_line else f"Lines {start}-{end_line}"
            )
            units.append(
                SourceUnit(cleaned, filename, location, category, len(units))
            )
        start = None
        buffered = []

    for line_number, line in enumerate(lines, start=1):
        if start is None:
            start = line_number
        buffered.append(line)
        buffered_chars = sum(len(item) for item in buffered)
        if len(buffered) >= 20 or buffered_chars >= 2000:
            emit(line_number)
    if buffered:
        emit(len(lines))
    if not units:
        raise DocumentExtractionError("The TXT file contains no usable text.")
    return units


def extract_document(
    file_bytes: bytes,
    filename: str,
    category: str,
) -> list[SourceUnit]:
    """Extract source-located units from a supported document."""

    if not file_bytes:
        raise EmptyDocumentError("The uploaded file is empty.")
    extension = Path(filename).suffix.lower()
    if extension not in _SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(_SUPPORTED_EXTENSIONS))
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{extension or '(none)'}'. Supported: {supported}."
        )
    extractors = {
        ".pdf": _extract_pdf,
        ".pptx": _extract_pptx,
        ".docx": _extract_docx,
        ".txt": _extract_txt,
    }
    return extractors[extension](file_bytes, filename, category)


def _make_chunk_id(unit: SourceUnit, normalized_text: str) -> str:
    identity = "\x1f".join(
        (unit.source, unit.location, unit.category, normalized_text)
    )
    return hashlib.sha256(identity.encode("utf-8")).hexdigest()


def chunk_source_units(
    units: list[SourceUnit],
    chunk_size_tokens: int,
    overlap_tokens: int,
) -> list[DocumentChunk]:
    """Split each source unit independently using token boundaries and overlap."""

    if chunk_size_tokens <= 0:
        raise ValueError("chunk_size_tokens must be positive.")
    if not 0 <= overlap_tokens < chunk_size_tokens:
        raise ValueError(
            "overlap_tokens must be at least 0 and smaller than chunk_size_tokens."
        )

    chunks: list[DocumentChunk] = []
    seen_ids: set[str] = set()
    for unit in units:
        tokens = encode_text(_clean_text(unit.text))
        start = 0
        unit_chunk_index = 0
        while start < len(tokens):
            end = min(start + chunk_size_tokens, len(tokens))
            chunk_text = _clean_text(decode_tokens(tokens[start:end]))
            if chunk_text:
                normalized = " ".join(chunk_text.split())
                chunk_id = _make_chunk_id(unit, normalized)
                if chunk_id not in seen_ids:
                    chunks.append(
                        DocumentChunk(
                            chunk_id=chunk_id,
                            text=chunk_text,
                            source=unit.source,
                            location=unit.location,
                            category=unit.category,
                            unit_index=unit.unit_index,
                            chunk_index=unit_chunk_index,
                            token_count=count_tokens(chunk_text),
                        )
                    )
                    seen_ids.add(chunk_id)
                    unit_chunk_index += 1
            if end >= len(tokens):
                break
            start = end - overlap_tokens
    return chunks


def ingest_document(
    file_bytes: bytes,
    filename: str,
    category: str,
    chunk_size_tokens: int,
    overlap_tokens: int,
) -> tuple[str, list[SourceUnit], list[DocumentChunk]]:
    """Hash, extract, and chunk one uploaded document."""

    file_hash = compute_file_hash(file_bytes)
    units = extract_document(file_bytes, filename, category)
    chunks = chunk_source_units(units, chunk_size_tokens, overlap_tokens)
    if not chunks:
        raise EmptyDocumentError("The document produced no usable chunks.")
    return file_hash, units, chunks

